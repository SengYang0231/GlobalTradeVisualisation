"""
FIT2179 A2 — Dataset 4 Preprocessing (Fixed)
Source: DOSM file_20260420112649.xlsx (Appendix iv)
Input files (place in same folder as this script):
  - file_20260420112649.xlsx
Outputs:
  - data/malaysia_exports_treemap.csv  → Chart 9 (treemap)

Fix applied: sector headers live in row[0], not row[1].
Old logic (checking row[1].isupper()) incorrectly labelled
all sub-sectors as "Other". Now correctly detects sector headers
from col 0 and sub-sector labels from col 1.
"""

import pandas as pd
import os
from openpyxl import load_workbook

DOSM_FILE  = "file_20260420112649.xlsx"
OUTPUT_DIR = "data"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Rows to skip entirely
SKIP_LABELS = {
    None, "", "Sector and Sub-sector",
    "Total Exports", "Table IV: Exports by Sector and Sub-sector ",
}


# ── Load Appendix iv ──────────────────────────────────────────────────────────
print("Loading DOSM Appendix iv — exports by sector...")
wb   = load_workbook(DOSM_FILE, read_only=True, data_only=True)
ws   = wb["Appendix iv"]
rows = list(ws.iter_rows(values_only=True))

records    = []
cur_sector = None

for row in rows:
    col0  = row[0]   # sector header lives here (e.g. "MANUFACTURING")
    col1  = row[1]   # sub-sector label lives here (e.g. "Electrical & Electronic Products")
    mar25 = row[2]
    feb26 = row[3]
    mar26 = row[4]
    share = row[5]

    # ── Sector header: col0 is uppercase text, col1 is None ──────────────────
    if col0 is not None and col1 is None:
        label = str(col0).strip()
        if label in SKIP_LABELS or not label:
            continue
        if label == label.upper() and len(label) > 2:
            cur_sector = label
            continue

    # ── Sub-sector row: col1 has the label, col0 is None ─────────────────────
    if col1 is not None and col0 is None:
        sub = str(col1).strip()
        if not sub or sub in SKIP_LABELS:
            continue
        if mar26 and isinstance(mar26, (int, float)):
            records.append({
                "sector":                  cur_sector or "Other",
                "sub_sector":              sub,
                "value_mar2026_rm_mil":    round(float(mar26), 2),
                "value_mar2025_rm_mil":    round(float(mar25), 2) if isinstance(mar25, (int, float)) else None,
                "value_feb2026_rm_mil":    round(float(feb26), 2) if isinstance(feb26, (int, float)) else None,
                "share_pct":               round(float(share), 4) if isinstance(share, (int, float)) else None,
            })

df = pd.DataFrame(records)
print(f"  Rows extracted: {len(df)}")
print(f"  Sectors found:  {list(df['sector'].unique())}")


# ════════════════════════════════════════════════════════════════════════════
# OUTPUT — malaysia_exports_treemap.csv  (Chart 9)
# One row per sub-sector, nested under sector for Vega treemap hierarchy
# ════════════════════════════════════════════════════════════════════════════
print("\nBuilding malaysia_exports_treemap.csv...")

treemap = df.dropna(subset=["value_mar2026_rm_mil"]).copy()
treemap = treemap.rename(columns={"value_mar2026_rm_mil": "value_rm_mil"})

# Hierarchy IDs for Vega treemap
treemap["id"]     = treemap["sector"] + " > " + treemap["sub_sector"]
treemap["parent"] = treemap["sector"]
treemap["label"]  = treemap["sub_sector"]

# Share calculations
total = treemap["value_rm_mil"].sum()
sector_totals = (
    treemap.groupby("sector")["value_rm_mil"]
    .sum()
    .rename("sector_total")
)
treemap = treemap.merge(sector_totals, on="sector")
treemap["share_total_pct"]     = (treemap["value_rm_mil"] / total * 100).round(2)
treemap["share_in_sector_pct"] = (treemap["value_rm_mil"] / treemap["sector_total"] * 100).round(2)

treemap = (
    treemap[[
        "sector", "sub_sector", "label", "id", "parent",
        "value_rm_mil", "share_total_pct", "share_in_sector_pct",
        "value_mar2025_rm_mil", "value_feb2026_rm_mil",
    ]]
    .sort_values(["sector", "value_rm_mil"], ascending=[True, False])
    .reset_index(drop=True)
)

treemap.to_csv(f"{OUTPUT_DIR}/malaysia_exports_treemap.csv", index=False)
print(f"  ✅ {len(treemap)} rows → {OUTPUT_DIR}/malaysia_exports_treemap.csv")

# Summary by sector
print("\n  Sector breakdown (RM million, Mar 2026):")
for sector, grp in treemap.groupby("sector"):
    s_total = grp["value_rm_mil"].sum()
    s_pct   = grp["share_total_pct"].sum()
    print(f"    {sector:<25}  {s_total:>10,.1f} RM mil  ({s_pct:.1f}% of exports)")
    for _, r in grp.nlargest(3, "value_rm_mil").iterrows():
        print(f"      ↳ {r['sub_sector']:<45} {r['value_rm_mil']:>9,.1f}")

print("\n✅ Dataset 4 complete.")
print("   malaysia_exports_treemap.csv → Chart 9")