"""
FIT2179 A2 — Dataset 2
Source: WITS Partner Timeseries (Malaysia → World row)
Outputs:
  - malaysia_trade_timeseries.csv   → Chart 4 (area chart, 1990–2023)
  - malaysia_trade_balance.csv      → Chart 5 (lollipop chart, 1990–2023)
"""

import pandas as pd
import os
from openpyxl import load_workbook

EXPORT_FILE = "WITS-Partner-Timeseries-export.xlsx"
IMPORT_FILE = "WITS-Partner-Timeseries-import.xlsx"
OUTPUT_DIR  = "data"
os.makedirs(OUTPUT_DIR, exist_ok=True)

MIN_YEAR = 1990
MAX_YEAR = 2023

EVENTS = {
    1997: "Asian Financial Crisis",
    2008: "Global Financial Crisis",
    2020: "COVID-19",
}


def extract_world_row(filepath, value_col):
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows         = list(ws.iter_rows(values_only=True))
    header       = rows[0]
    year_cols    = [c for c in header if str(c).isdigit()]
    year_indices = [list(header).index(y) for y in year_cols]
    for row in rows[1:]:
        partner = str(row[1]).strip() if row[1] else ""
        if partner == "World":
            return pd.DataFrame({
                "year":    [int(y) for y in year_cols],
                value_col: [row[i] for i in year_indices],
            })
    raise ValueError(f"'World' row not found in {filepath}")


print("Loading WITS files...")
exp_df = extract_world_row(EXPORT_FILE, "exports_usd_thou")
imp_df = extract_world_row(IMPORT_FILE, "imports_usd_thou")

mys = exp_df.merge(imp_df, on="year")
mys = mys[(mys["year"] >= MIN_YEAR) & (mys["year"] <= MAX_YEAR)].copy()

# Convert USD thousands → USD billions
mys["exports_usd_bn"] = (mys["exports_usd_thou"] / 1_000_000).round(2)
mys["imports_usd_bn"] = (mys["imports_usd_thou"] / 1_000_000).round(2)
mys["balance_usd_bn"] = (mys["exports_usd_bn"] - mys["imports_usd_bn"]).round(2)
mys["direction"]      = mys["balance_usd_bn"].apply(
    lambda x: "Surplus" if x >= 0 else "Deficit"
)
print(f"  Rows: {len(mys)} | Years: {mys['year'].min()}–{mys['year'].max()}")


# OUTPUT 1 — long format timeseries for Chart 4
print("\nBuilding malaysia_trade_timeseries.csv...")
ts_exp = mys[["year", "exports_usd_bn"]].rename(columns={"exports_usd_bn": "value_usd_bn"})
ts_exp["metric"] = "Exports"
ts_imp = mys[["year", "imports_usd_bn"]].rename(columns={"imports_usd_bn": "value_usd_bn"})
ts_imp["metric"] = "Imports"

timeseries = (
    pd.concat([ts_exp, ts_imp])
    .sort_values(["year", "metric"])
    .reset_index(drop=True)
)
timeseries["event"] = timeseries["year"].map(EVENTS).fillna("")
timeseries.to_csv(f"{OUTPUT_DIR}/malaysia_trade_timeseries.csv", index=False)
print(f"  Saved {len(timeseries)} rows → {OUTPUT_DIR}/malaysia_trade_timeseries.csv")
print(timeseries[timeseries["year"].isin([1997, 2008, 2020, 2023])].to_string(index=False))


# OUTPUT 2 — balance for Chart 5
print("\nBuilding malaysia_trade_balance.csv...")
balance = mys[["year", "exports_usd_bn", "imports_usd_bn",
               "balance_usd_bn", "direction"]].copy()
balance["abs_balance"] = balance["balance_usd_bn"].abs()
balance["event"]       = balance["year"].map(EVENTS).fillna("")
balance.to_csv(f"{OUTPUT_DIR}/malaysia_trade_balance.csv", index=False)
print(f"  Saved {len(balance)} rows → {OUTPUT_DIR}/malaysia_trade_balance.csv")
print(balance.tail(8).to_string(index=False))

print("\n✅ Dataset 2 complete.")
print("   malaysia_trade_timeseries.csv → Chart 4")
print("   malaysia_trade_balance.csv    → Chart 5")